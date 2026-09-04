"""Export canonical records to the 6-tab Google-Sheets-compatible format required
by the deliverables section: Startups / Products / Research Papers / Jobs / News /
Entity Mapping Log. Writes one CSV per tab plus a combined .xlsx workbook.
"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

TAB_HEADERS = {
    "Startups": ["schemaVersion", "recordType", "source.name", "source.url",
                 "content.entityName", "content.data.employeeCount", "collectedAt"],
    "Products": ["schemaVersion", "recordType", "source.name", "source.url",
                 "content.startupName", "content.pricingModel", "collectedAt"],
    "Research Papers": ["schemaVersion", "recordType", "content.title", "content.authors",
                         "content.paper_url", "content.github_url", "content.github_stars",
                         "content.published_date"],
    "Jobs": ["schemaVersion", "recordType", "content.company", "content.date",
             "content.is_remote", "content.role_family", "content.raw_title", "source.url"],
    "News": ["schemaVersion", "recordType", "content.title", "content.url",
             "content.published_at", "source.name"],
    "Entity Mapping Log": ["raw_name", "normalized_name", "canonical_name", "entity_type",
                            "match_method", "confidence", "source_url", "timestamp"],
}


def write_csv(rows: Iterable[dict], headers: list[str], out_path: str) -> None:
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def export_all_tabs(tabs: dict[str, list[dict]], out_dir: str = "data/exports") -> dict[str, str]:
    """tabs: {tab_name: [row_dict, ...]}. Returns {tab_name: csv_path}."""
    paths = {}
    for tab_name, rows in tabs.items():
        headers = TAB_HEADERS.get(tab_name, list(rows[0].keys()) if rows else [])
        path = f"{out_dir}/{tab_name.lower().replace(' ', '_')}.csv"
        write_csv(rows, headers, path)
        paths[tab_name] = path
    return paths


def export_xlsx(tabs: dict[str, list[dict]], out_path: str = "data/exports/final_submission.xlsx") -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for tab_name, rows in tabs.items():
        ws = wb.create_sheet(title=tab_name[:31])
        headers = TAB_HEADERS.get(tab_name, list(rows[0].keys()) if rows else [])
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
        for row in rows:
            ws.append([row.get(h, "") for h in headers])
        for i, header in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = max(14, min(40, len(header) + 4))
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
